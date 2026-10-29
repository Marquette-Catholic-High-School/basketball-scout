"""Roster file import (CSV / TSV / TXT / XLSX) with flexible header detection.

Coaches get rosters from MaxPreps, Hudl, a school website, or type them into Excel,
so this module is forgiving about layout and strict about the one thing film
matching needs: a clean jersey number for every player.

* Text files are decoded as UTF-8 (with or without a BOM) or UTF-16 with a BOM
  (Excel "Unicode Text"), falling back to Windows-1252 (Excel "CSV" on Windows).
  The delimiter (comma, semicolon, tab or pipe) is sniffed; when the sniffed one does
  not reveal the roster columns the others are tried in turn.
* Excel workbooks are read with ``read_only`` + ``data_only`` (formula results, low
  memory). Sheets named like "roster" are tried first, then the rest in order.
* Title rows above the header, blank rows, repeated header rows and unknown columns
  are skipped. Rows that cannot become a player are reported in ``skipped_rows``.
"""

from __future__ import annotations

import codecs
import csv
import io
import logging
import re
import warnings
from collections import Counter, defaultdict, deque
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from datetime import date
from pathlib import PurePath
from typing import Any, Literal

from openpyxl import load_workbook

from scout.errors import ErrorCode, ScoutError
from scout.schemas import Roster, RosterPlayer

logger = logging.getLogger(__name__)

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_PLAYERS = 200
MAX_TABLE_ROWS = 2000  # rows read per text file / worksheet; real rosters are far shorter
MAX_COLUMNS = 40
MAX_SHEETS = 10
HEADER_SCAN_ROWS = 20  # non-empty rows searched for the header row (title rows come first)
HEADERLESS_SCAN_ROWS = 50
HEADERLESS_MIN_SHARE = 0.6  # share of rows that must start with a jersey number in a header-less file
TEXT_EXTENSIONS = frozenset({".csv", ".tsv", ".txt"})
EXCEL_EXTENSIONS = frozenset({".xlsx"})
DELIMITERS = (",", "\t", ";", "|")

_ZIP_MAGIC = b"PK\x03\x04"  # every .xlsx is a zip archive
_OLE_MAGIC = b"\xd0\xcf\x11\xe0"  # legacy binary Office files (.xls)
_SAVE_AS_HINT = "In Excel use File > Save As and choose 'Excel Workbook (.xlsx)' or 'CSV'."

# Canonical field -> header spellings, compared after lowercasing and dropping everything
# except letters, digits and "#" ("Jersey #" -> "jersey#", "No." -> "no", "Gr." -> "gr").
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "jersey": (
        "#", "no", "nos", "nr", "num", "number", "nbr", "jersey", "jersey#", "jerseyno", "jerseynum",
        "jerseynumber", "uniform", "uniform#", "uniformno", "uniformnumber", "player#", "playerno", "playernumber",
    ),
    "name": (
        "name", "names", "player", "players", "playername", "fullname", "athlete", "athletename", "studentname",
        "lastfirst", "namelastfirst", "lastnamefirstname", "firstlast", "namefirstlast", "firstnamelastname",
    ),
    "first_name": ("first", "firstname", "fname", "givenname"),
    "last_name": ("last", "lastname", "lname", "surname", "familyname"),
    "suffix": ("suffix", "suf"),
    "position": ("pos", "position", "positions"),
    "grade": ("gr", "grade", "yr", "year", "class", "cl", "gradyear", "graduationyear", "classof"),
    "height": ("ht", "height", "hgt"),
    "notes": ("notes", "note", "comments", "comment", "remarks"),
    "fastscout_name": ("fastscoutname", "fastscout", "fsname", "fastscoutplayer", "fastscoutplayername"),
}
_HEADER_LOOKUP = {alias: canonical for canonical, aliases in COLUMN_ALIASES.items() for alias in aliases}
_NAME_FIELDS = ("name", "first_name", "last_name")
_FIELD_LABELS = {
    "jersey": "jersey", "name": "name", "first_name": "first name", "last_name": "last name",
    "position": "position", "grade": "grade", "height": "height",
}

_NAME_SUFFIXES = frozenset({"jr", "sr", "ii", "iii", "iv"})
_NAME_PARTICLES = frozenset({
    "van", "von", "der", "den", "de", "del", "della", "da", "di", "du", "la", "le",
    "st", "saint", "ter", "ten", "dos", "das", "bin", "al", "el",
})  # fmt: skip
_GRADES = {
    "Fr": ("fr", "fresh", "freshman", "9", "9th"),
    "So": ("so", "soph", "sophomore", "10", "10th"),
    "Jr": ("jr", "junior", "11", "11th"),
    "Sr": ("sr", "senior", "12", "12th"),
}
_GRADE_LOOKUP = {alias: label for label, aliases in _GRADES.items() for alias in aliases}
_POSITION_WORDS = frozenset({
    "g", "f", "c", "w", "p", "pg", "sg", "sf", "pf", "guard", "forward", "center", "centre", "post", "wing",
})  # fmt: skip

_JERSEY_CELL = re.compile(r"#?\s*(\d{1,2})(?:\.0+)?")  # "#5", "05", "00", "5.0" (Excel float)
_HEIGHT_VALUE = re.compile(r"\d\s*['’-]\s*\d{1,2}\s*(?:\"|”|'')?|\d{2}")  # 6-2, 6'2", 6’ 2”, 74
_EXCEL_DATE_HEIGHT = re.compile(r"(\d{1,2})-([a-z]{3})|([a-z]{3})-(\d{1,2})", re.IGNORECASE)  # "2-Jun" = 6-2
_EXCEL_TEXT_CELL = re.compile(r'\s*="(.*)"\s*')  # ="05": how some exports stop Excel dropping a leading zero
_HEIGHT_MONTHS = {"apr": 4, "may": 5, "jun": 6, "jul": 7}  # plausible feet values only
# Zero-width spaces / joiners and stray byte-order marks that rosters copied from web pages carry.
_INVISIBLE = dict.fromkeys(map(ord, "\u200b\u200c\u200d\u2060\ufeff"))

_Row = tuple[int, list[str]]  # (1-based row number in the file, tidied cell texts)


@dataclass
class RosterParseResult:
    players: list[RosterPlayer]
    # [{"row": 7, "reason": "..."}] with 1-based file rows (Excel row numbers for .xlsx)
    skipped_rows: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    detected_columns: dict[str, str] = field(default_factory=dict)  # canonical -> source header


@dataclass(frozen=True)
class _Layout:
    """Where the roster fields live in a table."""

    columns: dict[str, int]  # canonical field -> column index
    labels: dict[str, str]  # canonical field -> header text (or "column N" without a header)
    first_data_row: int  # index into the table's rows
    headerless: bool = False


def parse_roster_file(filename: str, content: bytes) -> RosterParseResult:
    """Dispatch on extension (.csv/.tsv/.txt -> text, .xlsx -> Excel).

    Header detection is case/space/punctuation-insensitive and finds the header row even
    when a title row precedes it. Recognised columns (examples):
    jersey: "#", "No", "No.", "Num", "Number", "Jersey", "Jersey #", "Uniform";
    name: "Name", "Player", "Player Name", "Full Name" (split "Last, First" and "First Last");
    first/last: "First", "First Name", "Last", "Last Name", "Surname";
    position: "Pos", "Position"; grade: "Gr", "Grade", "Yr", "Year", "Class", "Cl";
    height: "Ht", "Height" (keeps 6-2, 6'2", 74 as typed); notes: "Notes", "Comments";
    "FastScout Name" -> ``fastscout_name``.
    Also accepts header-less 2-3 column files whose first column is numeric.
    Jerseys are normalised (strip "#", spaces, ".0" from Excel floats) but "0" vs "00" is
    preserved. Grades "Sr"/"12"/"Senior" become "Sr" (likewise Fr/So/Jr); other grade
    values (e.g. a graduation year) are kept as typed. Duplicate jerseys -> warning, later
    row kept as separate player.
    Raises ``ScoutError(ROSTER_UNSUPPORTED_FILE)`` for other extensions and
    ``ScoutError(ROSTER_PARSE_FAILED)`` when no player rows are found, the file is larger
    than 2 MB, or it holds more than 200 players.
    """
    extension = PurePath(filename or "").suffix.lower()
    _check_file(extension, content)
    # A workbook renamed to .csv still starts with the zip signature: read it as Excel.
    if extension in EXCEL_EXTENSIONS or content.startswith(_ZIP_MAGIC):
        tables: Iterable[list[_Row]] = _workbook_tables(content)
    else:
        tables = _text_tables(_decode(content), tab_first=extension == ".tsv")
    result = _parse_tables(tables)
    logger.info("Roster file %r: %d players, %d rows skipped", filename, len(result.players), len(result.skipped_rows))
    return result


def merge_players(
    existing: Roster, incoming: list[RosterPlayer], mode: Literal["replace", "merge"] = "replace"
) -> Roster:
    """``replace``: incoming becomes the roster. ``merge``: update players with the same
    jersey (keeping their ids so track overrides stay valid), append new jerseys.

    In both modes an incoming player whose jersey is already on the roster takes over
    the existing player's id, so manual track assignments survive a re-import. ``merge``
    only overwrites fields the incoming row actually fills (a file without a FastScout
    column keeps the FastScout names typed earlier). ``existing`` is not modified.
    """
    if mode == "replace":
        unused_ids: dict[str, deque[str]] = defaultdict(deque)
        for player in existing.players:
            unused_ids[player.jersey].append(player.id)
        players = []
        for new in incoming:
            ids = unused_ids[new.jersey]
            players.append(new.model_copy(update={"id": ids.popleft()}) if ids else new)
    elif mode == "merge":
        players = list(existing.players)
        open_slots: dict[str, deque[int]] = defaultdict(deque)
        for index, player in enumerate(players):
            open_slots[player.jersey].append(index)
        for new in incoming:
            if open_slots[new.jersey]:
                index = open_slots[new.jersey].popleft()
                players[index] = _updated_player(players[index], new)
            else:
                players.append(new)
    else:
        raise ValueError(f"unknown roster import mode {mode!r} (expected 'replace' or 'merge')")
    return existing.model_copy(update={"players": players})


def split_name(full_name: str) -> tuple[str, str]:
    """Split one name cell into ``(first, last)``.

    With a comma it is "Last, First" ("Van Dyke, Kevin", "Smith Jr., John"), unless the
    comma only sets off a suffix ("John Smith, Jr."). Without one it is "First Last" and
    the last name keeps particles ("Kevin Van Dyke") and suffixes ("John Smith Jr.",
    "Wilson III"). A single word is taken as the last name.
    """
    text = " ".join(full_name.split())
    if "," not in text:
        return _split_first_last(text)
    parts = [part.strip() for part in text.split(",") if part.strip()]
    if not parts:
        return "", ""
    head, rest = parts[0], parts[1:]
    suffixes = [part for part in rest if _is_suffix(part)]
    given = [part for part in rest if not _is_suffix(part)]
    if not given:  # "John Smith, Jr."
        first, last = _split_first_last(head)
        return first, " ".join([last, *suffixes])
    first_tokens = " ".join(given).split()
    trailing: list[str] = []
    while len(first_tokens) > 1 and _is_suffix(first_tokens[-1]):  # "Smith, John Jr."
        trailing.insert(0, first_tokens.pop())
    return " ".join(first_tokens), " ".join([head, *suffixes, *trailing])


# ---------------------------------------------------------------------------
# File checks and reading
# ---------------------------------------------------------------------------


def _check_file(extension: str, content: bytes) -> None:
    """Reject files we cannot or should not read, with a coach-friendly reason."""
    if extension not in TEXT_EXTENSIONS | EXCEL_EXTENSIONS:
        excel_like = extension in {".xls", ".xlsm", ".xlsb", ".ods", ".numbers"}
        raise ScoutError(
            ErrorCode.ROSTER_UNSUPPORTED_FILE,
            f"Roster files must be .csv, .tsv, .txt or .xlsx (this one is '{extension or 'no extension'}').",
            hint=_SAVE_AS_HINT if excel_like else None,
        )
    if not content.strip():
        raise ScoutError(ErrorCode.ROSTER_PARSE_FAILED, "The roster file is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise ScoutError(
            ErrorCode.ROSTER_PARSE_FAILED,
            f"The roster file is {len(content) / 1_048_576:.1f} MB; roster files are limited to 2 MB.",
            hint="A roster only needs number and name columns. Save just the roster sheet as CSV and import that.",
        )
    if content.startswith(_OLE_MAGIC):
        raise ScoutError(
            ErrorCode.ROSTER_UNSUPPORTED_FILE, "This is an old-style Excel (.xls) file.", hint=_SAVE_AS_HINT
        )


def _decode(content: bytes) -> str:
    """Bytes -> text: BOM-marked UTF-16/UTF-8 first, then UTF-8, then Windows-1252."""
    if content.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        text = content.decode("utf-16", errors="replace")
    else:
        if content.startswith(codecs.BOM_UTF8):
            content = content[len(codecs.BOM_UTF8) :]
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            # Excel's plain "CSV" on Windows writes cp1252 (accents in names like Peña).
            text = content.decode("cp1252", errors="replace")
    return text.replace("\x00", "")


def _text_tables(text: str, *, tab_first: bool) -> Iterator[list[_Row]]:
    """The text split with each plausible delimiter, the most likely one first."""
    for delimiter in _delimiter_candidates(text, tab_first=tab_first):
        yield _read_delimited(text, delimiter)


def _delimiter_candidates(text: str, *, tab_first: bool) -> list[str]:
    """csv.Sniffer's guess first (tab for .tsv), then every other supported delimiter."""
    sample = "\n".join(text.splitlines()[:HEADERLESS_SCAN_ROWS])
    try:
        sniffed = [csv.Sniffer().sniff(sample, delimiters="".join(DELIMITERS)).delimiter]
    except csv.Error:  # title rows or a single column confuse the sniffer
        sniffed = []
    preferred = ["\t"] if tab_first else []
    return list(dict.fromkeys([*preferred, *sniffed, *DELIMITERS]))


def _read_delimited(text: str, delimiter: str) -> list[_Row]:
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, skipinitialspace=True)
    rows: list[_Row] = []
    try:
        for number, cells in enumerate(reader, start=1):
            if number > MAX_TABLE_ROWS:
                logger.warning("Roster file has more than %d rows; ignoring the rest", MAX_TABLE_ROWS)
                break
            tidy = _tidy(_unwrap_excel_text(cell) for cell in cells)
            if tidy:
                rows.append((number, tidy))
    except csv.Error as exc:
        raise ScoutError(
            ErrorCode.ROSTER_PARSE_FAILED, "The roster file is not a readable CSV file.", detail=str(exc)
        ) from exc
    return rows


def _unwrap_excel_text(cell: str) -> str:
    """'="05"' -> '05' (a CSV cell written as an Excel text formula keeps "05" and "00" as typed)."""
    match = _EXCEL_TEXT_CELL.fullmatch(cell)
    return match.group(1) if match else cell


def _workbook_tables(content: bytes) -> list[list[_Row]]:
    """Rows of each worksheet, sheets named like "roster" first, empty sheets dropped."""
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")  # openpyxl warns about Excel features it does not read
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                sheets = sorted(workbook.worksheets, key=lambda sheet: "roster" not in sheet.title.lower())
                tables = [_sheet_rows(sheet) for sheet in sheets[:MAX_SHEETS]]
            finally:
                workbook.close()
    except Exception as exc:  # openpyxl raises many types for damaged, encrypted or fake .xlsx files
        raise ScoutError(
            ErrorCode.ROSTER_PARSE_FAILED,
            "The Excel file could not be opened. It may be damaged, password-protected, or not really an .xlsx file.",
            hint="Open it in Excel and save it again as .xlsx or CSV.",
            detail=f"{type(exc).__name__}: {exc}",
        ) from exc
    return [rows for rows in tables if rows]


def _sheet_rows(sheet: Any) -> list[_Row]:
    rows: list[_Row] = []
    # Explicit bounds: some exporters write wrong sheet dimensions, and stray formatting
    # can extend a sheet to row 1,048,576.
    for number, cells in enumerate(sheet.iter_rows(max_row=MAX_TABLE_ROWS, max_col=MAX_COLUMNS), start=1):
        tidy = _tidy(_excel_text(cell.value, cell.number_format) for cell in cells)
        if tidy:
            rows.append((number, tidy))
    return rows


def _excel_text(value: Any, number_format: str | None) -> str:
    """An Excel cell value as the text the coach sees in the cell."""
    if value is None:
        return ""
    if isinstance(value, date):
        # Excel silently turns a height typed as 6-2 into the date 2-Jun; month-day gives
        # the height back. No roster column we read holds real dates.
        return f"{value.month}-{value.day}"
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int | float):
        if isinstance(value, float) and not value.is_integer():
            return f"{value:g}"
        # A "00" number format is how Excel shows jersey 00 (or 05) stored as a number.
        width = len(number_format) if number_format and set(number_format) == {"0"} else 0
        return f"{int(value):0{width}d}"
    return str(value)


def _tidy(cells: Iterable[Any]) -> list[str]:
    """Collapse whitespace inside every cell (dropping invisible characters) and drop trailing empty cells."""
    tidy = [" ".join(str(cell).translate(_INVISIBLE).split()) for cell in cells][:MAX_COLUMNS]
    while tidy and not tidy[-1]:
        tidy.pop()
    return tidy


# ---------------------------------------------------------------------------
# Finding the columns
# ---------------------------------------------------------------------------


def _parse_tables(tables: Iterable[list[_Row]]) -> RosterParseResult:
    """Players from the first table (worksheet or delimiter variant) that has a roster."""
    saw_rows = False
    columns_without_players = False
    for rows in tables:
        if not rows:
            continue
        saw_rows = True
        layout = _find_header(rows) or _headerless_layout(rows)
        if layout is None:
            continue
        result = _read_players(rows, layout)
        if result.players:
            return result
        columns_without_players = True
    if not saw_rows:
        raise ScoutError(ErrorCode.ROSTER_PARSE_FAILED, "The roster file is empty.")
    if columns_without_players:
        raise ScoutError(
            ErrorCode.ROSTER_PARSE_FAILED, "Found the roster columns, but no rows under them have a jersey number."
        )
    raise ScoutError(ErrorCode.ROSTER_PARSE_FAILED, "Could not find the jersey number and name columns in this file.")


def _header_key(cell: str) -> str:
    return re.sub(r"[^a-z0-9#]", "", cell.lower())


def _header_columns(cells: list[str]) -> dict[str, int]:
    """Canonical field -> column index for every recognised header cell (first one wins)."""
    columns: dict[str, int] = {}
    for index, cell in enumerate(cells):
        canonical = _HEADER_LOOKUP.get(_header_key(cell))
        if canonical and canonical not in columns:
            columns[canonical] = index
    return columns


def _find_header(rows: list[_Row]) -> _Layout | None:
    """The first row near the top that names a jersey column and a name column."""
    for index, (_, cells) in enumerate(rows[:HEADER_SCAN_ROWS]):
        columns = _header_columns(cells)
        if "jersey" in columns and any(name in columns for name in _NAME_FIELDS):
            labels = {canonical: cells[column] for canonical, column in columns.items()}
            return _Layout(columns=columns, labels=labels, first_data_row=index + 1)
    return None


def _headerless_layout(rows: list[_Row]) -> _Layout | None:
    """Files without a header: 2-3 columns, most rows starting with a jersey number.

    The second column is the name. A third column is read as position, grade or height
    when most of its values look like one, or as the last name when it holds words
    ("10, John, Smith").
    """
    sample = [cells for _, cells in rows[:HEADERLESS_SCAN_ROWS]]
    numbered = [cells for cells in sample if 2 <= len(cells) <= 3 and _clean_jersey(cells[0])]
    if not numbered or len(numbered) < HEADERLESS_MIN_SHARE * len(sample):
        return None
    columns = {"jersey": 0}
    third = _column_kind([cells[2] for cells in numbered if len(cells) == 3])
    if third == "name":
        columns.update(first_name=1, last_name=2)
    else:
        columns["name"] = 1
        if third:
            columns[third] = 2
    labels = {canonical: f"column {column + 1}" for canonical, column in columns.items()}
    return _Layout(columns=columns, labels=labels, first_data_row=0, headerless=True)


def _column_kind(values: list[str]) -> str | None:
    """The kind most values in a header-less column share, if more than half agree."""
    if not values:
        return None
    kind, count = Counter(_value_kind(value) for value in values).most_common(1)[0]
    return kind if count * 2 > len(values) else None


def _value_kind(value: str) -> str | None:
    parts = [part for part in re.split(r"[/\-\s]+", value.lower()) if part]
    if parts and all(part in _POSITION_WORDS for part in parts):
        return "position"
    if _grade_key(value) in _GRADE_LOOKUP or re.fullmatch(r"20\d\d", value):
        return "grade"
    if _HEIGHT_VALUE.fullmatch(value):
        return "height"
    if any(char.isalpha() for char in value):
        return "name"
    return None


# ---------------------------------------------------------------------------
# Rows -> players
# ---------------------------------------------------------------------------


def _read_players(rows: list[_Row], layout: _Layout) -> RosterParseResult:
    result = RosterParseResult(players=[], detected_columns=dict(layout.labels))
    if layout.headerless:
        read_as = ", ".join(_FIELD_LABELS[canonical] for canonical in layout.columns)
        result.warnings.append(f"No header row was found, so the columns were read as: {read_as}.")
    rows_by_number: dict[str, list[tuple[int, str]]] = defaultdict(list)  # jersey key -> [(row, jersey)]
    for number, cells in rows[layout.first_data_row :]:
        values = {canonical: _cell(cells, column) for canonical, column in layout.columns.items()}
        raw_jersey = values.get("jersey", "")
        if _HEADER_LOOKUP.get(_header_key(raw_jersey)) == "jersey":
            continue  # the header row repeated (multi-page exports)
        first, last = _player_names(values)
        jersey = _clean_jersey(raw_jersey)
        if jersey is None:
            reason = f"invalid jersey number {raw_jersey!r}" if raw_jersey else "missing jersey number"
            result.skipped_rows.append({"row": number, "reason": reason})
            if raw_jersey and (first or last):
                full_name = " ".join(part for part in (first, last) if part)
                result.warnings.append(
                    f"Row {number}: {raw_jersey!r} is not a jersey number (0-99), so {full_name} was skipped."
                )
            continue
        if not (first or last):
            result.warnings.append(f"Row {number}: #{jersey} has no name.")
        result.players.append(_make_player(jersey, first, last, values))
        if len(result.players) > MAX_PLAYERS:
            raise ScoutError(
                ErrorCode.ROSTER_PARSE_FAILED,
                f"The file has more than {MAX_PLAYERS} player rows, which is not a single team's roster.",
                hint="Check that this is the right file, or split it into one roster per team.",
            )
        rows_by_number[_jersey_key(jersey)].append((number, jersey))
    result.warnings.extend(_duplicate_warnings(rows_by_number))
    return result


def _cell(cells: list[str], column: int) -> str:
    return cells[column] if column < len(cells) else ""


def _make_player(jersey: str, first: str, last: str, values: dict[str, str]) -> RosterPlayer:
    return RosterPlayer(
        jersey=jersey,
        first_name=first,
        last_name=last,
        position=values.get("position") or None,
        grade=_clean_grade(values.get("grade", "")),
        height=_clean_height(values.get("height", "")),
        notes=values.get("notes") or None,
        fastscout_name=values.get("fastscout_name") or None,
    )


def _player_names(values: dict[str, str]) -> tuple[str, str]:
    """(first, last) from First/Last columns when filled, else from the name column."""
    first, last = values.get("first_name", ""), values.get("last_name", "")
    if not (first or last):
        first, last = split_name(values.get("name", ""))
    suffix = values.get("suffix", "")
    if suffix and not last.endswith(suffix):
        last = f"{last} {suffix}".strip()
    return first, last


def _split_first_last(text: str) -> tuple[str, str]:
    tokens = text.split()
    suffixes: list[str] = []
    while len(tokens) > 1 and _is_suffix(tokens[-1]):
        suffixes.insert(0, tokens.pop())
    if len(tokens) <= 1:
        return "", " ".join(tokens + suffixes)
    start = len(tokens) - 1  # the last name starts at the final word, plus any particles before it
    while start > 1 and tokens[start - 1].lower().strip(".") in _NAME_PARTICLES:
        start -= 1
    return " ".join(tokens[:start]), " ".join(tokens[start:] + suffixes)


def _is_suffix(word: str) -> bool:
    return word.lower().strip(".") in _NAME_SUFFIXES


def _clean_jersey(raw: str) -> str | None:
    """'#5' -> '5', '5.0' -> '5'; '05', '0' and '00' stay as typed; anything else -> None."""
    match = _JERSEY_CELL.fullmatch(raw.strip())
    return match.group(1) if match else None


def _jersey_key(jersey: str) -> str:
    """Jerseys film matching cannot tell apart ("5" and "05"); "0" and "00" stay distinct."""
    return jersey.lstrip("0") or jersey


def _grade_key(raw: str) -> str:
    key = re.sub(r"[^a-z0-9]", "", raw.lower())  # "Sr." -> "sr", "Gr 12" -> "gr12"
    return key.removeprefix("grade").removeprefix("gr")


def _clean_grade(raw: str) -> str | None:
    """'Sr', 'Sr.', '12', '12th', 'Senior' -> 'Sr' (same for Fr/So/Jr); other values as typed."""
    return _GRADE_LOOKUP.get(_grade_key(raw)) or raw or None


def _clean_height(raw: str) -> str | None:
    """Heights stay as typed (6-2, 6'2", 74), except Excel's date version of 6-2 ("2-Jun")."""
    match = _EXCEL_DATE_HEIGHT.fullmatch(raw)
    if match:
        inches = int(match.group(1) or match.group(4))
        feet = _HEIGHT_MONTHS.get((match.group(2) or match.group(3)).lower())
        if feet and inches <= 11:
            return f"{feet}-{inches}"
    return raw or None


def _updated_player(old: RosterPlayer, new: RosterPlayer) -> RosterPlayer:
    """``old`` (same id and jersey) with every field ``new`` fills in."""
    fields = ("first_name", "last_name", "position", "grade", "height", "notes", "fastscout_name")
    return old.model_copy(update={name: value for name in fields if (value := getattr(new, name))})


def _duplicate_warnings(rows_by_number: dict[str, list[tuple[int, str]]]) -> list[str]:
    messages: list[str] = []
    for entries in rows_by_number.values():
        if len(entries) < 2:
            continue
        jerseys = " / ".join(f"#{jersey}" for jersey in dict.fromkeys(jersey for _, jersey in entries))
        rows = ", ".join(str(row) for row, _ in entries)
        messages.append(
            f"Jersey {jerseys} is listed {len(entries)} times (rows {rows}). All were kept; "
            "remove the extra so film matching is not ambiguous."
        )
    return messages
